#!/usr/bin/env python3
"""构建大麦演示静态发行版（实施计划 Task 4；确定性，可重复构建零漂移）。

用法（仓库根或 services/api 均可）：
    python3 scripts/build_damai_demo.py --output fixtures/damai
    python3 scripts/build_damai_demo.py --output fixtures/damai --check

产物（规格 §4.2）：
    canonical/*.jsonl、workbooks/damai_logistics_full_v1.xlsx、
    statements/damai_fy2025.yaml、damai_fy2026.yaml、
    forecast/rolling_forecast.jsonl、manifest.json、README.md
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

import yaml

REPOSITORY_ROOT = Path(__file__).resolve().parent.parent
if str(REPOSITORY_ROOT / "services/api/src") not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT / "services/api/src"))

from flow_api.data_contract.contract import load_contract
from flow_api.data_contract.workbook import render_workbook
from flow_api.fixtures.damai.canonical import (
    build_damai_canonical_package,
)
from flow_api.fixtures.damai.generator import build_damai_package
from flow_api.fixtures.damai.operations import build_damai_operations_payloads
from flow_api.fixtures.damai.statements import (
    build_damai_statement_payloads,
)
from flow_api.fixtures.generator import write_canonical_package

CONTRACT_PATH = REPOSITORY_ROOT / "templates/excel/flow_v1_contract.yaml"
GENERATOR_VERSION = "damai-demo-v1"
_WORKBOOK_REL = "workbooks/damai_logistics_full_v1.xlsx"
_STATEMENT_RELS = {
    "FY2025": "statements/damai_fy2025.yaml",
    "FY2026": "statements/damai_fy2026.yaml",
}
_FORECAST_REL = "forecast/rolling_forecast.jsonl"
_OPERATIONS_RELS = (
    "operations/damai_segment_series.yaml",
    "operations/damai_operating_metrics.yaml",
)
_MANIFEST_REL = "manifest.json"
_README_REL = "README.md"

README_TEXT = """---
doc_id: FLOW-GEN-DAMAI-FIXTURE-README-001
title: 大麦物流演示发行版 fixtures 说明
doc_type: generated
status: generated
version: 1.0
created_at: 2026-09-24
updated_at: 2026-09-24
owner: FLOW
generator_ref: scripts/build_damai_demo.py
input_hash: deterministic-static
---

# 大麦物流演示发行版（synthetic）

本目录由 `scripts/build_damai_demo.py` 确定性生成，**禁止手工修改**；重建命令：

    python3 scripts/build_damai_demo.py --output fixtures/damai --check

- `canonical/*.jsonl`：flow.excel.v1 canonical 数据包（批次/期间/维度/事实）；
- `workbooks/damai_logistics_full_v1.xlsx`：标准工作簿（Finance BP 导入用，
  actual/budget，无 forecast——预测在 sidecar，`persistence: static-only`、
  `page_coverage: excluded`，不得冒充已上线能力）；
- `statements/damai_fy2025.yaml`、`damai_fy2026.yaml`：闭合合成财报（六大恒等锚）；
- `operations/damai_segment_series.yaml`、`damai_operating_metrics.yaml`：
  DAMAI.SYN 独立分部序列与运营事实（synthetic，血缘指向 manifest.json）；
- `forecast/rolling_forecast.jsonl`：静态预测 sidecar；
- `manifest.json`：期间、行数、核心汇总、逐文件 SHA-256、`synthetic: true`。

装载：`python3 scripts/seed_damai_demo.py`（幂等，经 IntakeService/领域服务链，
见规格《大麦物流演示数据设计》§4.3）。
"""

REQUIRED_FILES: tuple[str, ...] = (
    "canonical/ar_collections.jsonl",
    "canonical/batch.jsonl",
    "canonical/customer_segments.jsonl",
    "canonical/customers.jsonl",
    "canonical/financial_actuals.jsonl",
    "canonical/logistics_products.jsonl",
    "canonical/management_accounts.jsonl",
    "canonical/monthly_budgets.jsonl",
    "canonical/operating_actuals.jsonl",
    "canonical/organizations.jsonl",
    "canonical/periods.jsonl",
    "canonical/regions.jsonl",
    "canonical/scenario_versions.jsonl",
    _WORKBOOK_REL,
    _STATEMENT_RELS["FY2025"],
    _STATEMENT_RELS["FY2026"],
    _FORECAST_REL,
    *_OPERATIONS_RELS,
    _MANIFEST_REL,
    _README_REL,
)


def _workbook_semantic_digest(path: Path) -> str:
    """xlsx 的单元格值级指纹：zip 容器（时间戳/顺序）不参与，语义漂移才漂移。"""

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    payload: list[Any] = []
    for sheet in workbook.worksheets:
        payload.append([sheet.title])
        for row in sheet.iter_rows(values_only=True):
            payload.append(["" if v is None else str(v) for v in row])
    workbook.close()
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    return "semantic:" + hashlib.sha256(body.encode("utf-8")).hexdigest()


def _sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _decimal(value: Any) -> Any:
    from decimal import Decimal

    if isinstance(value, Decimal):
        return format(value, ".4f")
    return value


def _dump_statements(
    destination: Path, package: dict[str, Any]
) -> dict[str, int]:
    payloads = build_damai_statement_payloads(package)
    counts: dict[str, int] = {}
    for fy, rel in _STATEMENT_RELS.items():
        payload = payloads[fy]
        target = destination / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            yaml.safe_dump(payload, allow_unicode=True, sort_keys=True, width=100),
            encoding="utf-8",
        )
        statement_rows = payload.get("statements") or payload.get("report") or payload
        counts[rel] = len(statement_rows) if isinstance(statement_rows, (list, dict)) else 0
    return counts


def _dump_operations(destination: Path, package: dict[str, Any]) -> None:
    payloads = build_damai_operations_payloads(package)
    for key, rel in zip(
        ("segment_series", "operating_metrics"), _OPERATIONS_RELS, strict=True
    ):
        target = destination / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            yaml.safe_dump(payloads[key], allow_unicode=True, sort_keys=True, width=100),
            encoding="utf-8",
        )


def _dump_forecast(destination: Path, sidecar: dict[str, Any]) -> None:
    target = destination / _FORECAST_REL
    target.parent.mkdir(parents=True, exist_ok=True)
    rows = sidecar.get("rows", [])
    lines = [
        json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        for row in rows
    ]
    target.write_bytes(("\n".join(lines) + ("\n" if lines else "")).encode("utf-8"))


def _summaries(destination: Path) -> dict[str, str]:
    from decimal import Decimal

    total_revenue = Decimal(0)
    path = destination / "canonical/operating_actuals.jsonl"
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            total_revenue += Decimal(json.loads(line)["revenue"])
    return {"operating_revenue_total": format(total_revenue, ".4f")}


def build_release(destination: Path) -> dict[str, Any]:
    """构建完整发行版并返回 {manifest, files: {rel: sha256}}。"""

    destination.mkdir(parents=True, exist_ok=True)
    raw = build_damai_package()
    package = build_damai_canonical_package()

    canonical_dir = destination / "canonical"
    write_canonical_package(package, canonical_dir)

    contract = load_contract(CONTRACT_PATH)
    render_workbook(contract, package, destination / _WORKBOOK_REL)

    statement_counts = _dump_statements(destination, raw)
    _dump_forecast(destination, raw["forecast_sidecar"])
    _dump_operations(destination, raw)

    (destination / _README_REL).write_text(README_TEXT, encoding="utf-8")

    written = tuple(rel for rel in REQUIRED_FILES if rel != _MANIFEST_REL)
    files: dict[str, str] = {}
    for rel in written:
        target = destination / rel
        if not target.is_file():
            raise FileNotFoundError(f"发行版缺失文件：{rel}")
        # xlsx 的 zip 容器非字节确定（openpyxl 时间戳/顺序），以语义指纹入册
        files[rel] = (
            _workbook_semantic_digest(target) if rel == _WORKBOOK_REL else _sha256_of(target)
        )

    row_counts = {
        rel: sum(
            1
            for line in (destination / rel).read_text(encoding="utf-8").splitlines()
            if line.strip()
        )
        for rel in written
        if rel.startswith("canonical/") and rel.endswith(".jsonl")
    }
    manifest: dict[str, Any] = {
        "generator": "scripts/build_damai_demo.py",
        "generator_version": GENERATOR_VERSION,
        "synthetic": True,
        "profile_version": raw["profile_version"],
        "batch_code": package.batch.batch_code,
        "periods": {
            "total": len(package.periods),
            "comparison": sum(1 for p in package.periods if p.window == "comparison"),
            "analysis": sum(1 for p in package.periods if p.window == "analysis"),
            "analysis_start": package.batch.analysis_start_month,
            "analysis_end": package.batch.analysis_end_month,
        },
        "row_counts": row_counts,
        "statement_row_counts": statement_counts,
        "summaries": _summaries(destination),
        "files": {rel: files[rel] for rel in written},
    }
    (destination / _MANIFEST_REL).write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    files[_MANIFEST_REL] = _sha256_of(destination / _MANIFEST_REL)
    return {"manifest": manifest, "files": files}


def check_release(destination: Path) -> int:
    with __import__("tempfile").TemporaryDirectory() as td:
        rebuilt = build_release(Path(td))
    committed = {
        rel: (
            _workbook_semantic_digest(destination / rel)
            if rel == _WORKBOOK_REL
            else _sha256_of(destination / rel)
        )
        for rel in REQUIRED_FILES
    }
    drift = [rel for rel in REQUIRED_FILES if committed[rel] != rebuilt["files"][rel]]
    for rel in drift:
        print(f"DRIFT {rel}: committed {committed[rel]} != rebuilt {rebuilt['files'][rel]}", file=sys.stderr)
    print("release check: " + ("PASS" if not drift else "FAIL"))
    return 0 if not drift else 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=REPOSITORY_ROOT / "fixtures/damai")
    parser.add_argument("--check", action="store_true", help="重建对账已提交产物")
    args = parser.parse_args(argv)
    root = args.output.resolve()
    if args.check:
        return check_release(root)
    result = build_release(root)
    print(f"release built at {root} ({len(result['files'])} files)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
