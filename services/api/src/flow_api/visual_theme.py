"""将可移植 FLOW 视觉 token 映射到文档导出表面。"""
from __future__ import annotations

import platform
from typing import Any

from flow_api.visual_tokens import VISUAL_TOKENS


def _primary_font(token: str) -> str:
    native_fonts = {
        "Darwin": {
            "typography-family-display": "Songti SC",
            "typography-family-body": "PingFang SC",
        },
        "Windows": {
            "typography-family-display": "SimSun",
            "typography-family-body": "Microsoft YaHei",
        },
        "Linux": {
            "typography-family-display": "Noto Serif CJK SC",
            "typography-family-body": "Noto Sans CJK SC",
        },
    }
    native = native_fonts.get(platform.system(), {}).get(token)
    if native:
        return native
    return VISUAL_TOKENS[token].split(",", maxsplit=1)[0].strip().strip("\"'")


def _rgb(token: str) -> Any:
    from pptx.dml.color import RGBColor

    value = VISUAL_TOKENS[token].removeprefix("#")
    return RGBColor.from_string(value)  # type: ignore[no-untyped-call]


def apply_pptx_visual_theme(presentation: Any) -> None:
    """统一文稿纸面底色、标题字级和古金识别线，不改报告数据与分页结构。"""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    title_color = _rgb("colors-brand-navy")
    body_color = _rgb("colors-text-primary")
    paper_color = _rgb("colors-surface-canvas")
    gold_color = _rgb("colors-brand-gold")
    for slide in presentation.slides:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = paper_color
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, Inches(0.075)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = gold_color
        accent.line.fill.background()
        title_shape = slide.shapes.title
        for shape in list(slide.shapes):
            if shape is accent:
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            is_title = title_shape is not None and (
                shape is title_shape or shape.name == title_shape.name
            )
            family = _primary_font(
                "typography-family-display" if is_title else "typography-family-body"
            )
            color = title_color if is_title else body_color
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.name = family
                paragraph.font.color.rgb = color
                if is_title:
                    paragraph.font.size = Pt(26)
                    paragraph.font.bold = True
                for run in paragraph.runs:
                    run.font.name = family
                    run.font.color.rgb = color
                    if is_title:
                        run.font.size = Pt(26)
                        run.font.bold = True


def style_workbook_headers(workbook: Any) -> None:
    """统一 XLSX 表头、数据阅读和打印设置，保留工作簿的数据优先结构。"""
    from copy import copy

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    fill = PatternFill("solid", fgColor=VISUAL_TOKENS["colors-surface-nav"].lstrip("#"))
    font = Font(
        name=_primary_font("typography-family-body"),
        bold=True,
        color=VISUAL_TOKENS["colors-text-inverse"].lstrip("#"),
        size=10,
    )
    body_font = _primary_font("typography-family-body")
    body_color = VISUAL_TOKENS["colors-text-primary"].lstrip("#")
    zebra = PatternFill("solid", fgColor=VISUAL_TOKENS["colors-surface-canvas"].lstrip("#"))
    line = Side(style="hair", color=VISUAL_TOKENS["colors-line-subtle"].lstrip("#"))
    for sheet in workbook.worksheets:
        if sheet.max_row < 1:
            continue
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = VISUAL_TOKENS["colors-brand-gold"].lstrip("#")
        if sheet.sheet_properties.pageSetUpPr is None:
            sheet.sheet_properties.pageSetUpPr = PageSetupProperties(
                fitToPage=True, autoPageBreaks=False
            )
        else:
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape" if sheet.max_column >= 6 else "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1 if sheet.max_column <= 9 else 2
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.3
        sheet.page_margins.right = 0.3
        sheet.page_margins.top = 0.45
        sheet.page_margins.bottom = 0.45
        sheet.page_margins.header = 0.2
        sheet.page_margins.footer = 0.2
        sheet.print_title_rows = "1:1"
        sheet.oddFooter.center.text = "FLOW · 第 &P 页 / 共 &N 页"
        sheet.oddFooter.center.size = 8
        sheet.oddFooter.center.font = "Noto Sans SC,Regular"
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = copy(cell.font)
                cell.font = Font(
                    name=body_font,
                    size=10,
                    bold=cell.font.bold,
                    italic=cell.font.italic,
                    color=body_color,
                )
                cell.alignment = Alignment(
                    horizontal="right" if isinstance(cell.value, (int, float)) else "left",
                    vertical="top",
                    wrap_text=True,
                )
                cell.border = Border(bottom=line)
                if cell.row % 2 == 0:
                    cell.fill = zebra
        for column in sheet.columns:
            first_cell = column[0]
            letter = first_cell.column_letter
            current_width = sheet.column_dimensions[letter].width
            if current_width is not None and current_width != 13:
                continue
            width = min(
                max(len(str(cell.value or "")) for cell in column) + 2,
                42,
            )
            sheet.column_dimensions[letter].width = max(width, 10)
        if sheet.max_column == 1:
            sheet.column_dimensions["A"].width = max(
                sheet.column_dimensions["A"].width or 0, 84
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
