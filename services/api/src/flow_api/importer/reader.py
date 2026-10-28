"""xlsx 读取（R2）：把上传的模板文件解析为校验器输入行。

只认中文表头，按列名取值；表头不符时报缺失列错误。
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from flow_api.importer.templates import (
    BANK_STATEMENT_COLUMNS,
    BANK_STATEMENT_SHEET,
    INCOME_STATEMENT_COLUMNS,
    INCOME_STATEMENT_SHEET,
    TRIAL_BALANCE_COLUMNS,
    TRIAL_BALANCE_SHEET,
)


class SheetReadError(ValueError):
    """表头缺失或工作表不存在。"""


def _read_sheet(data: bytes, sheet_title: str, columns: tuple[str, ...]) -> list[dict[str, str]]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        if sheet_title not in workbook.sheetnames:
            raise SheetReadError(f"缺少工作表「{sheet_title}」")
        ws = workbook[sheet_title]
        iterator = ws.iter_rows(values_only=True)
        header_row: tuple[Any, ...] | None = None
        for candidate in iterator:
            if any(value is not None and str(value).strip() for value in candidate):
                header_row = candidate
                break
        if header_row is None:
            raise SheetReadError(f"工作表「{sheet_title}」为空")
        header_normalized = [str(value or "").strip() for value in header_row]
        missing = [column for column in columns if column not in header_normalized]
        if missing:
            raise SheetReadError(f"缺少必需列：{'、'.join(missing)}")
        col_index = {column: header_normalized.index(column) for column in columns}

        rows: list[dict[str, str]] = []
        for values in iterator:
            if values is None or all(value is None or str(value).strip() == "" for value in values):
                continue
            row: dict[str, str] = {}
            for column in columns:
                raw = values[col_index[column]]
                row[column] = "" if raw is None else str(raw).strip()
            rows.append(row)
        return rows
    finally:
        workbook.close()


def read_trial_balance(data: bytes) -> list[dict[str, str]]:
    return _read_sheet(data, TRIAL_BALANCE_SHEET, TRIAL_BALANCE_COLUMNS)


def read_income_statement(data: bytes) -> list[dict[str, str]] | None:
    workbook = load_workbook(BytesIO(data), read_only=True)
    has_sheet = INCOME_STATEMENT_SHEET in workbook.sheetnames
    workbook.close()
    if not has_sheet:
        return None  # 利润表可选
    return _read_sheet(data, INCOME_STATEMENT_SHEET, INCOME_STATEMENT_COLUMNS)


def read_bank_statement(data: bytes) -> list[dict[str, str]] | None:
    workbook = load_workbook(BytesIO(data), read_only=True)
    has_sheet = BANK_STATEMENT_SHEET in workbook.sheetnames
    workbook.close()
    if not has_sheet:
        return None  # 银行流水可选（M3）
    return _read_sheet(data, BANK_STATEMENT_SHEET, BANK_STATEMENT_COLUMNS)
