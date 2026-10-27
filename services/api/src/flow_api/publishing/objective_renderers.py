"""从冻结载荷（objective.v1）渲染 HTML / XLSX 报告。

数据边界：渲染只消费冻结 payload（不可变），任何数值都是载荷原值的
格式化投影；不重算、不补造（D043/D049）。PDF 复用既有 Chromium 打印链。
"""

from __future__ import annotations

import io
import re
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from flow_api.visual_theme import style_workbook_headers
from flow_api.visual_tokens import VISUAL_TOKENS

_GOLDEN_SCALE_NOTE = "单位以 unit_note 为准，数值为披露原值的字符串投影"


def _source_lines(payload: dict[str, Any]) -> list[str]:
    source = payload.get("source", {})
    return [
        f"公司：{source.get('company_name', '')}（{source.get('stock_code', '')}）",
        f"期间：{source.get('period_label', '')} {source.get('report_kind', '')}",
        f"来源：{source.get('source_ref', '')}",
        f"单位：{source.get('unit_note', '')}",
    ]


def _statements(payload: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    return cast(dict[str, list[dict[str, Any]]], payload.get("statements", {}))


_VALUE_FIELDS = ("value_end", "value_begin", "value_current", "value_prior")


def _row_values(row: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(row.get(field) for field in _VALUE_FIELDS)


def _reconciliation(payload: dict[str, Any]) -> list[tuple[str, str]]:
    """归一行 vs 原始行逐表值集合核对：归一化只改名不改变披露值。

    核对只输出「一致/不一致」判定，不产生任何新财务数字（D043/D049）。
    """

    from collections import Counter

    normalized = _statements(payload)
    raw = cast(dict[str, list[dict[str, Any]]], payload.get("statements_raw", {}))
    results: list[tuple[str, str]] = []
    for statement_type in sorted(set(normalized) | set(raw)):
        n_values = Counter(_row_values(row) for row in normalized.get(statement_type, []))
        r_values = Counter(_row_values(row) for row in raw.get(statement_type, []))
        verdict = "一致" if n_values == r_values else "不一致"
        results.append((statement_type, verdict))
    return results


def render_html_from_payload(payload: dict[str, Any]) -> str:
    """极简 HTML：来源身份头 + 四表逐行原值（上期列标注比较基准）+ 勾稽核对 + 口径与溯源。

    数据边界：渲染只消费冻结 payload（不可变），任何数值都是载荷原值的
    格式化投影；不重算、不补造（D043/D049）。PDF 复用既有 Chromium 打印链。
    """

    blocks = []
    source = payload.get("source", {})
    source_sha256 = str(source.get("source_sha256", ""))
    source_href_prefix = (
        f"/api/v1/statements/sources/{source_sha256}/content"
        if source.get("source_available") is True
        and re.fullmatch(r"[0-9a-f]{64}", source_sha256)
        else None
    )
    for statement_type, rows in _statements(payload).items():
        rendered_rows: list[str] = []
        for row in rows:
            page_number = row.get("page_number")
            source_link = (
                f" <a href='{source_href_prefix}#page={page_number}'>原文 p{page_number} ↗</a>"
                if source_href_prefix and isinstance(page_number, int) and page_number > 0
                else ""
            )
            rendered_rows.append(
                f"<tr><td>{row['item']}{source_link}</td>"
                f"<td class='num'>{row.get('value_current') or ''}</td>"
                f"<td class='num'>{row.get('value_prior') or ''}</td></tr>"
            )
        row_html = "".join(rendered_rows)
        blocks.append(
            f"<h3>{statement_type}</h3><table>"
            "<tr><th>项目</th><th>本期/期末</th><th>上期/期初（比较基准）</th></tr>"
            f"{row_html}</table>"
        )

    recon_rows = "".join(
        f"<tr><td>{statement_type}</td>"
        f"<td class='{'ok' if verdict == '一致' else 'bad'}'>{verdict}</td></tr>"
        for statement_type, verdict in _reconciliation(payload)
    )
    frozen_at = payload.get("frozen_at", "")
    src_block = "".join(
        f"<div>{line}</div>" for line in _source_lines(payload)
    )
    sha256 = str(source.get("source_sha256", ""))
    # 叙事结构（借鉴 #9 客观部分）：现状判断 → 勾稽与质量核对 → 口径与溯源。
    # 原因推断与行动章不在客观报告范围（U5 主观证据门禁），显式标注。
    html_doc = (
        "<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width, initial-scale=1'>"
        "<title>客观财报分析报告</title>"
        f"<style>@page{{size:A4;margin:14mm}}*{{box-sizing:border-box}}"
        f"body{{font-family:{VISUAL_TOKENS['typography-family-body']};margin:32px;"
        f"color:{VISUAL_TOKENS['colors-text-primary']};"
        f"background:{VISUAL_TOKENS['colors-surface-canvas']};line-height:1.55}}"
        f"h1,h2,h3{{font-family:{VISUAL_TOKENS['typography-family-display']};"
        f"color:{VISUAL_TOKENS['colors-brand-navy']}}}"
        f"h1{{font-size:26px;padding-bottom:12px;"
        f"border-bottom:2px solid {VISUAL_TOKENS['colors-brand-gold']}}}"
        f"h2{{font-size:18px;border-left:3px solid {VISUAL_TOKENS['colors-brand-gold']};"
        "padding-left:10px}"
        "table{border-collapse:collapse;width:100%;margin:12px 0}"
        f"td,th{{border-bottom:1px solid {VISUAL_TOKENS['colors-line-default']};"
        f"padding:7px 10px;text-align:left;vertical-align:top}}"
        f"th{{color:{VISUAL_TOKENS['colors-brand-navy']};font-weight:600;"
        f"background:{VISUAL_TOKENS['colors-surface-soft']}}}"
        f".num{{text-align:right;font-variant-numeric:tabular-nums;"
        f"font-family:{VISUAL_TOKENS['typography-family-numeric']}}}"
        "h2{margin-top:24px}h3{font-size:14px;margin-top:16px}"
        f".src{{color:{VISUAL_TOKENS['colors-text-muted']};font-size:13px;margin-bottom:18px}}"
        f".ok{{color:{VISUAL_TOKENS['colors-semantic-positive']}}}"
        f".bad{{color:{VISUAL_TOKENS['colors-semantic-negative']};font-weight:600}}"
        f".note{{color:{VISUAL_TOKENS['colors-text-muted']};font-size:12px}}"
        f".chapter{{margin-top:28px;border-top:2px solid {VISUAL_TOKENS['colors-line-default']};"
        f"padding-top:8px}}.chapter-label{{color:{VISUAL_TOKENS['colors-brand-goldStrong']};"
        "font-size:12px;font-weight:600;letter-spacing:0.1em}"
        f".scope{{border:1px solid {VISUAL_TOKENS['colors-line-strong']};"
        f"border-left:3px solid {VISUAL_TOKENS['colors-brand-gold']};"
        f"background:{VISUAL_TOKENS['colors-surface-paper']};padding:12px 16px;"
        f"color:{VISUAL_TOKENS['colors-text-secondary']}}}"
        f"@media(max-width:700px){{body{{margin:16px}}td,th{{padding:6px}}}}"
        f"@media print{{body{{margin:0}}tr{{break-inside:avoid}}h1,h2,h3{{break-after:avoid}}}}"
        "</style></head><body>"
        "<h1>客观财报分析报告</h1>"
        f"<div class='src'>{src_block}冻结时间：{frozen_at}</div>"
        "<div class='chapter'><span class='chapter-label'>第一章 · 现状判断</span>"
        "<h2>披露原值（比较基准：上期/期初）</h2></div>"
        f"{''.join(blocks)}"
        "<div class='chapter'><span class='chapter-label'>第二章 · 勾稽与质量核对</span>"
        "<h2>归一化 vs 原始披露</h2></div>"
        f"<table><tr><th>报表</th><th>核对结果</th></tr>{recon_rows}</table>"
        "<p class='note'>核对仅验证归一化未改变披露值；本报告不产生新财务数字，"
        "所有数值均为冻结载荷原值的投影。差额分解仅解释金额构成，"
        "不证明业务原因（C14）。</p>"
        "<div class='chapter'><span class='chapter-label'>第三章 · 口径与溯源</span>"
        "<h2>口径声明</h2></div>"
        f"<div class='src'>单位口径：{source.get('unit_note', '')}<br/>"
        f"比较基准：上期/期初为披露同期对照值<br/>"
        f"来源文件：{source.get('source_ref', '')}<br/>"
        f"来源 SHA-256：<code>{sha256[:16]}</code>…（完整值见冻结载荷）</div>"
        "<div class='scope'><strong>范围说明：</strong>本报告为客观事实层"
        "（现状判断与质量核对）。原因推断、改善行动章<b>不在客观报告范围</b>"
        "——主观内容须经证据门槛审批后方可进入独立报告（U5 门禁，"
        "确定性原语仅输出提示复核信号，不断言因果）。</div>"
        "</body></html>"
    )
    return html_doc


def render_xlsx_bytes_from_payload(payload: dict[str, Any]) -> bytes:
    """XLSX：口径与证据页（含来源指纹与勾稽核对）+ 各报表逐行原值（金额列右对齐）。"""

    wb = Workbook()
    evidence = cast(Worksheet, wb.active)
    evidence.title = "口径与证据"
    evidence.column_dimensions["A"].width = 74
    evidence.column_dimensions["B"].width = 18
    evidence["A1"] = "FLOW 客观财报分析（冻结载荷投影）"
    evidence["A1"].font = Font(bold=True)
    for idx, line in enumerate(_source_lines(payload), start=3):
        evidence[f"A{idx}"] = line
    source = payload.get("source", {})
    sha256 = str(source.get("source_sha256", ""))
    row_cursor = len(_source_lines(payload)) + 4
    evidence[f"A{row_cursor}"] = f"来源 SHA-256：{sha256}"
    row_cursor += 1
    evidence[f"A{row_cursor}"] = f"冻结时间：{payload.get('frozen_at', '')}"
    row_cursor += 1
    evidence[f"A{row_cursor}"] = "比较基准：上期/期初为披露同期对照值"
    row_cursor += 2
    evidence[f"A{row_cursor}"] = "勾稽核对（归一化 vs 原始披露）："
    evidence[f"A{row_cursor}"].font = Font(bold=True)
    for statement_type, verdict in _reconciliation(payload):
        row_cursor += 1
        evidence[f"A{row_cursor}"] = statement_type
        evidence[f"B{row_cursor}"] = verdict
    row_cursor += 2
    evidence[f"A{row_cursor}"] = _GOLDEN_SCALE_NOTE

    for statement_type, rows in _statements(payload).items():
        ws = wb.create_sheet(statement_type[:28])
        ws.append(["项目", "本期/期末", "上期/期初（比较基准）"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row["item"], row.get("value_current"), row.get("value_prior")])
        for column, width in (("A", 42), ("B", 20), ("C", 26)):
            ws.column_dimensions[column].width = width
        for data_row in ws.iter_rows(min_row=2):
            for cell in data_row[1:]:
                cell.alignment = Alignment(horizontal="right")

    buffer = io.BytesIO()
    style_workbook_headers(wb)
    wb.save(buffer)
    return buffer.getvalue()
