"""经营轨六主题概览 HTML 渲染（O3）。

数据边界与 objective_renderers 相同：只消费 typed 概览（引擎已算好的
确定性结果），不重算、不补造；复用 objective 的 Chromium 打印链。
叙事按借鉴 #13 客观部分组织：现状判断 → 质量核对 → 口径与溯源；
行动章显式标注不在客观报告范围（U5 门禁，借鉴 #9）。
"""

from __future__ import annotations

from html import escape
from typing import Any

from flow_api.operations.engine import OperationsOverview

_DIRECTION_LABEL = {"negative": "关注", "warning": "提示复核"}
_ASSURANCE_LABEL = {
    "audited": "经审计",
    "unaudited": "未经审计",
    "management_disclosure": "管理层披露",
}


def _validated_overview(
    overview: OperationsOverview | dict[str, Any],
) -> OperationsOverview:
    if isinstance(overview, OperationsOverview):
        return overview
    return OperationsOverview.model_validate(overview["overview"])


def render_operations_html(overview: OperationsOverview | dict[str, Any]) -> str:
    """接受 typed 概览或冻结载荷 dict（schema 同构）。"""

    overview = _validated_overview(overview)
    blocks: list[str] = []
    for theme in overview.themes:
        if theme.status == "not_applicable":
            blocks.append(
                f"<section class='theme na'><h2>{theme.name}</h2>"
                f"<p class='muted'>待内部数据（{theme.reason}）——"
                "披露缺失与内部数据缺口如实标注，不推测填补。</p></section>"
            )
            continue
        rows = []
        for metric in theme.metrics:
            if metric.status == "computed":
                caliber = (
                    f"<div class='caliber'>口径：{metric.caliber_note}</div>"
                    if metric.caliber_note
                    else ""
                )
                evidence = ""
                if metric.source == "operating_fact":
                    assurance = _ASSURANCE_LABEL.get(metric.assurance, metric.assurance or "未标注")
                    page = f" · 第 {escape(metric.source_page)} 页" if metric.source_page else ""
                    evidence = (
                        "<div class='evidence'>"
                        f"期间：{escape(metric.period_label)} · {escape(assurance)} · "
                        f"来源：{escape(metric.source_ref)}{page} · "
                        f"SHA-256：{escape(metric.source_sha256[:16])}…"
                        "</div>"
                    )
                rows.append(
                    f"<tr><td>{metric.name}</td>"
                    f"<td class='num'>{metric.value}</td>"
                    f"<td>{metric.basis or '—'}{caliber}{evidence}</td></tr>"
                )
            else:
                rows.append(
                    f"<tr><td>{metric.name}</td>"
                    f"<td class='muted'>暂不可算（{metric.reason}）</td>"
                    f"<td class='muted'>缺失不补造</td></tr>"
                )
        blocks.append(
            f"<section class='theme'><h2>{theme.name}</h2>"
            f"<table><tr><th>指标</th><th>值</th><th>比较基准与口径</th></tr>"
            f"{''.join(rows)}</table></section>"
        )

    watch_items = "".join(
        f"<li data-direction='{item['direction']}'>"
        f"<span class='tag'>{_DIRECTION_LABEL.get(item['direction'], item['direction'])}</span>"
        f"{item['message']}</li>"
        for item in overview.management_watch
    )
    watch_block = (
        f"<h2>管理关注（≤3 条，提示复核）</h2><ul class='watch'>{watch_items}</ul>"
        if overview.management_watch
        else "<h2>管理关注</h2><p class='muted'>本期无确定性提示信号。</p>"
    )

    return (
        "<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>"
        "<title>经营分析概览</title>"
        "<style>body{font-family:'PingFang SC',sans-serif;margin:32px;color:#1e293b}"
        "table{border-collapse:collapse;width:100%;margin:12px 0}"
        "td,th{border-bottom:1px solid #eee;padding:6px 10px;text-align:left}"
        "th{color:#475569;font-weight:600}"
        ".num{text-align:right;font-variant-numeric:tabular-nums}"
        "h1{font-size:22px}h2{font-size:16px;margin-top:24px}"
        ".muted{color:#64748b;font-size:13px}"
        ".caliber{color:#64748b;font-size:12px}"
        ".evidence{color:#64748b;font-size:11px;margin-top:3px;overflow-wrap:anywhere}"
        ".watch{list-style:none;padding:0;display:flex;flex-direction:column;gap:8px}"
        ".watch li{border:1px solid #e2e8f0;border-left-width:4px;border-radius:6px;"
        "padding:8px 12px;font-size:14px}"
        ".watch li[data-direction='warning']{border-left-color:#d97706}"
        ".watch li[data-direction='negative']{border-left-color:#b91c1c}"
        ".tag{display:inline-block;margin-right:8px;font-size:12px;font-weight:600;"
        "color:#475569}"
        ".scope{border:1px solid #cbd5e1;border-radius:6px;padding:10px 14px;"
        "color:#475569;font-size:13px;margin-top:24px}</style></head><body>"
        "<h1>经营分析概览</h1>"
        f"<div class='muted'>目录：{overview.catalog_id}（L1 结果层，"
        "数据可得性分层见 D050）</div>"
        + "".join(blocks)
        + watch_block
        + "<div class='scope'><strong>范围说明：</strong>本报告只承载客观事实与"
        "可复算的确定性信号；原因推断与<b>行动</b>建议<b>不在客观报告范围</b>"
        "（主观内容需证据门槛，见 U5 门禁）。</div>"
        "</body></html>"
    )


def render_operations_xlsx(
    overview: OperationsOverview | dict[str, Any],
) -> bytes:
    """从同一冻结载荷生成可复核 XLSX；值保持字符串精度，不在渲染层重算。"""

    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    typed = _validated_overview(overview)
    workbook = Workbook()
    facts = workbook.active
    assert facts is not None
    facts.title = "六主题概览"
    facts.append(
        [
            "主题",
            "指标",
            "状态",
            "值",
            "比较基准",
            "口径",
            "期间",
            "保障状态",
            "来源",
            "原文页码",
            "来源 SHA-256",
            "不可用原因",
        ]
    )
    for theme in typed.themes:
        if not theme.metrics:
            facts.append(
                [theme.name, "", theme.status, "", "", "", "", "", "", "", "", theme.reason or ""]
            )
        for metric in theme.metrics:
            facts.append(
                [
                    theme.name,
                    metric.name,
                    metric.status,
                    metric.value or "",
                    metric.basis or "",
                    metric.caliber_note or "",
                    metric.period_label,
                    _ASSURANCE_LABEL.get(metric.assurance, metric.assurance),
                    metric.source_ref,
                    metric.source_page,
                    metric.source_sha256,
                    metric.reason or "",
                ]
            )
    facts.freeze_panes = "A2"
    facts.auto_filter.ref = facts.dimensions

    watch = workbook.create_sheet("管理关注")
    watch.append(["代码", "方向", "事实陈述"])
    for item in typed.management_watch:
        watch.append(
            [
                item["code"],
                _DIRECTION_LABEL.get(item["direction"], item["direction"]),
                item["message"],
            ]
        )

    trace = workbook.create_sheet("版本与溯源")
    trace.append(["项目", "值"])
    trace.append(["报告 ID", typed.report_id])
    trace.append(["分析目录", typed.catalog_id])
    trace.append(["范围边界", "原因推断与行动建议不在客观报告范围"])
    trace.append(["渲染原则", "只消费冻结载荷；缺失不补造；不在渲染层重算"])

    for sheet in workbook.worksheets:
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
            column_index = column[0].column
            if isinstance(column_index, int):
                sheet.column_dimensions[get_column_letter(column_index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_operations_pptx(
    overview: OperationsOverview | dict[str, Any],
) -> bytes:
    """从冻结载荷生成管理层 PPTX，保留值/基准/口径/来源四要素。"""

    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches, Pt

    typed = _validated_overview(overview)
    presentation = Presentation()

    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = "经营分析概览"
    cover.placeholders[
        1
    ].text = f"六主题客观事实报告\n目录 {typed.catalog_id}\n报告 ID {typed.report_id}"

    def add_slide(title: str, lines: list[str]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        frame = slide.shapes.add_textbox(
            Inches(0.55), Inches(1.35), Inches(12.2), Inches(5.7)
        ).text_frame
        frame.word_wrap = True
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(13)
            paragraph.space_after = Pt(7)

    watch_lines = [
        f"{_DIRECTION_LABEL.get(item['direction'], item['direction'])}｜{item['message']}"
        for item in typed.management_watch
    ]
    add_slide("管理关注", watch_lines or ["本期无确定性提示信号。"])

    for theme in typed.themes:
        lines: list[str] = []
        if theme.status == "not_applicable":
            lines.append(f"待内部数据（{theme.reason}）——缺失不补造。")
        for metric in theme.metrics:
            if metric.status != "computed":
                lines.append(f"{metric.name}｜暂不可算（{metric.reason}）")
                continue
            lines.append(
                f"{metric.name}｜{metric.value}｜基准：{metric.basis or '—'}｜"
                f"口径：{metric.caliber_note or '—'}"
            )
            if metric.source == "operating_fact":
                assurance = _ASSURANCE_LABEL.get(metric.assurance, metric.assurance)
                lines.append(
                    f"  证据：{metric.period_label} · {assurance} · {metric.source_ref}"
                    f"{f' · 第 {metric.source_page} 页' if metric.source_page else ''} · "
                    f"SHA-256 {metric.source_sha256}"
                )
        add_slide(theme.name, lines or ["本主题暂无可展示条目。"])

    add_slide(
        "范围与使用边界",
        [
            "本报告只承载客观事实与可复算的确定性信号。",
            "原因推断与行动建议不在客观报告范围。",
            "所有格式均消费同一冻结载荷；缺失不补造；渲染层不重算。",
        ],
    )

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


__all__ = [
    "render_operations_html",
    "render_operations_pptx",
    "render_operations_xlsx",
]
