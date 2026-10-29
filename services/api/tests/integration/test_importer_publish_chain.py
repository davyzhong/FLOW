"""R5 九页接线集成测试：导入账套 → 治理发布链 → 驾驶舱出数 → 删除回落。

验收口径（work_item FLOW-WI-REAL-DATA-IMPORT-20261028-001 R5）：
- 口径等价：trial_balance_v1 目录逐期间快照值 == 独立复算期望值；
  流量指标（收入/费用/净利润）全部期间快照之和 == derive_metrics 直读全期合计；
- 治理切换：发布后 DashboardService 总览以导入快照为活跃源，revenue 卡出真数，
  科目余额表推不出的指标（orders/gross_margin 等）如实 unavailable；
- 真实门禁：勾稽不平（资产−负债−权益≠0）被发布门禁拒绝；
- 删除级联：删除数据集下线整条发布链，仪表盘回到删除前状态。
"""

from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from flow_api.api.routes.importer import set_store
from flow_api.dashboard.models import ActiveFilters
from flow_api.dashboard.repositories import DashboardSourceRepository
from flow_api.dashboard.service import DashboardService
from flow_api.importer.normalization import derive_metrics
from flow_api.importer.store import ImportStore
from flow_api.infrastructure.db import get_engine
from flow_api.infrastructure.models.analytics import (
    AnalysisRun,
    MetricDefinition,
    MetricSnapshot,
    MetricValue,
)
from flow_api.infrastructure.models.intake import AnalysisBatch
from flow_api.main import create_app

_HEADERS = {"Authorization": "Bearer flow-dev-bp"}
_PERIODS: tuple[str, ...] = ("2025-01", "2025-02")
# 每行：科目编码/名称/类别/方向/期初/借方发生/贷方发生/期末。
# 期末资产 20000→21000、负债 8000→7000、权益 12000→14000 → 每期勾稽均为 0；
# 损益明细（收入/成本/销售/管理）与利润表样例逐期一致（净利润勾稽通过）。
_ROWS: dict[str, tuple[tuple[str, str, str, str, str, str, str, str], ...]] = {
    "2025-01": (
        ("1001", "库存现金", "资产", "借", "10000", "0", "0", "12000"),
        ("1002", "银行存款", "资产", "借", "10000", "0", "0", "8000"),
        ("2001", "短期借款", "负债", "贷", "10000", "0", "0", "8000"),
        ("3001", "实收资本", "权益", "贷", "10000", "0", "0", "12000"),
        ("6001", "主营业务收入", "损益", "贷", "0", "0", "5000", "0"),
        ("6401", "主营业务成本", "损益", "借", "0", "1800", "0", "0"),
        ("6601", "销售费用", "损益", "借", "0", "200", "0", "0"),
        ("6602", "管理费用", "损益", "借", "0", "1000", "0", "0"),
    ),
    "2025-02": (
        ("1001", "库存现金", "资产", "借", "12000", "0", "0", "13000"),
        ("1002", "银行存款", "资产", "借", "8000", "0", "0", "8000"),
        ("2001", "短期借款", "负债", "贷", "8000", "0", "0", "7000"),
        ("3001", "实收资本", "权益", "贷", "12000", "0", "0", "14000"),
        ("6001", "主营业务收入", "损益", "贷", "0", "0", "7000", "0"),
        ("6401", "主营业务成本", "损益", "借", "0", "2400", "0", "0"),
        ("6601", "销售费用", "损益", "借", "0", "300", "0", "0"),
        ("6602", "管理费用", "损益", "借", "0", "1300", "0", "0"),
    ),
}
# 利润表样例（与 _ROWS 损益发生额逐期一致；净利润勾稽通过）
_INCOME_ROWS: dict[str, tuple[tuple[str, str], ...]] = {
    "2025-01": (
        ("营业收入", "5000"),
        ("营业成本", "1800"),
        ("销售费用", "200"),
        ("管理费用", "1000"),
        ("净利润", "2000"),
    ),
    "2025-02": (
        ("营业收入", "7000"),
        ("营业成本", "2400"),
        ("销售费用", "300"),
        ("管理费用", "1300"),
        ("净利润", "3000"),
    ),
}
# 银行流水样例（M3；余额连续：期初锚 100000 → 112000 → 104000 → …）
_BANK_ROWS: dict[str, tuple[tuple[str, str, str, str], ...]] = {
    "2025-01": (
        ("2025-01-05", "收到货款", "12000", "0", "112000"),
        ("2025-01-18", "支付房租", "0", "8000", "104000"),
    ),
    "2025-02": (
        ("2025-02-06", "收到货款", "15000", "0", "119000"),
        ("2025-02-20", "支付工资", "0", "6000", "113000"),
    ),
}
_TB_HEADER = (
    "会计期间",
    "科目编码",
    "科目名称",
    "科目类别",
    "方向",
    "期初余额",
    "本期借方发生额",
    "本期贷方发生额",
    "期末余额",
)
_IS_HEADER = ("会计期间", "项目", "本期金额", "本年累计金额")
_BANK_HEADER = ("交易日期", "摘要", "收入金额", "支出金额", "账户余额")


def _expected_period_metrics(period: str) -> dict[str, Decimal]:
    """从样例行独立复算七项指标（不依赖被测代码的期望值）。"""
    totals = {"资产": Decimal("0"), "负债": Decimal("0"), "权益": Decimal("0")}
    revenue = expenses = Decimal("0")
    for code, _name, category, _direction, _opening, debit, credit, closing in _ROWS[period]:
        amount_debit = Decimal(debit)
        amount_credit = Decimal(credit)
        amount_closing = Decimal(closing)
        if category in totals:
            totals[category] += amount_closing
        if code.startswith("6"):  # 损益类样例科目（6001 收入 / 6602 费用）
            revenue += amount_credit
            expenses += amount_debit
    assets, liabilities, equity = totals["资产"], totals["负债"], totals["权益"]
    return {
        "revenue": revenue,
        "total_expenses": expenses,
        "net_profit": revenue - expenses,
        "total_assets": assets,
        "total_liabilities": liabilities,
        "total_equity": equity,
        # 引擎 output_scale=6（trial_balance_v1 目录声明）
        "debt_ratio": (liabilities / assets).quantize(Decimal("0.000001")),
    }


def _expected_income_metrics(period: str) -> dict[str, Decimal]:
    """M2 利润表指标期望：独立复算营业成本/毛利/毛利率/费用分项。

    rd/finance 样例无数据 → 缺席（返回集中不出现）。
    """
    items = {item: Decimal(amount) for item, amount in _INCOME_ROWS[period]}
    revenue = _expected_period_metrics(period)["revenue"]
    cost = items["营业成本"]
    return {
        "operating_cost": cost,
        "gross_profit": revenue - cost,
        "gross_margin": ((revenue - cost) / revenue).quantize(Decimal("0.000001")),
        "selling_expense": items["销售费用"],
        "admin_expense": items["管理费用"],
    }


def _expected_cash_metrics(period: str) -> dict[str, Decimal]:
    """M3 现金口径期望：银行流水按期间聚合（收付实现制）。"""
    inflow = outflow = Decimal("0")
    for _date, _summary, income, expense, _balance in _BANK_ROWS[period]:
        inflow += Decimal(income)
        outflow += Decimal(expense)
    return {
        "cash_inflow": inflow,
        "cash_outflow": outflow,
        "net_cash_flow": inflow - outflow,
    }


def _expected_detail_metrics(period: str) -> dict[str, Decimal]:
    """M1 结构指标期望：样例科目按前缀映射组独立复算（组内方向净额）。

    样例只有 1001/1002（货币资金组）命中——其余组无命中科目，
    事实缺席 → 指标缺席（不写 0 冒充真实余额）。
    """
    cash = Decimal("0")
    for code, _name, category, direction, _opening, _debit, _credit, closing in _ROWS[period]:
        if code.startswith(("1001", "1002")) and category == "资产":
            cash += Decimal(closing) if direction == "借" else -Decimal(closing)
    return {"cash": cash}


def _snapshot_values(session: Session, snapshot_id: str) -> dict[str, Decimal]:
    """无维度、actual_month 口径的快照值（与总览 total scope 匹配粒度一致）。"""
    rows = session.execute(
        select(MetricDefinition.metric_code, MetricValue.exact_value)
        .join(MetricDefinition, MetricDefinition.id == MetricValue.metric_definition_id)
        .where(
            MetricValue.metric_snapshot_id == UUID(snapshot_id),
            MetricValue.comparison_type == "actual_month",
            MetricValue.organization_id.is_(None),
        )
    ).all()
    return {code: Decimal(value) for code, value in rows}


def _workbook_bytes(
    *,
    with_income: bool = False,
    net_override: str | None = None,
    with_bank: bool = False,
    break_bank_balance: bool = False,
) -> bytes:
    """样例工作簿；可选附利润表（net_override 用于勾稽拒绝）与银行流水。"""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "科目余额表"
    ws.append(list(_TB_HEADER))
    for period in _PERIODS:
        for row in _ROWS[period]:
            ws.append([period, *row])
    if with_income:
        ws_income = workbook.create_sheet("利润表")
        ws_income.append(list(_IS_HEADER))
        for period in _PERIODS:
            for item, amount in _INCOME_ROWS[period]:
                if net_override is not None and item == "净利润":
                    amount = net_override
                ws_income.append([period, item, amount, amount])
    if with_bank:
        ws_bank = workbook.create_sheet("银行流水")
        ws_bank.append(list(_BANK_HEADER))
        for period in _PERIODS:
            for date_text, summary, income, expense, balance in _BANK_ROWS[period]:
                if break_bank_balance and summary == "支付房租":
                    balance = "999999"
                ws_bank.append([date_text, summary, income, expense, balance])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _identity_breaking_bytes() -> bytes:
    """借贷平衡但勾稽不平：借方合计=贷方合计，而资产−负债−权益=−200。"""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "科目余额表"
    ws.append(list(_TB_HEADER))
    ws.append(["2025-03", "1001", "库存现金", "资产", "借", "100", "0", "0", "100"])
    ws.append(["2025-03", "3001", "实收资本", "权益", "借", "100", "0", "0", "100"])
    ws.append(["2025-03", "3002", "资本公积", "权益", "贷", "200", "0", "0", "200"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def chain_env() -> Iterator[tuple[TestClient, Session]]:
    """同一 engine 上的 ImportStore + 治理 Session + TestClient。"""
    engine = get_engine()
    set_store(ImportStore(engine))
    session = Session(engine, expire_on_commit=False)
    client = TestClient(create_app())
    yield client, session
    session.rollback()
    session.close()


def _upload(client: TestClient, content: bytes, name: str) -> str:
    response = client.post(
        f"/api/v1/importer/dataset?name={name}",
        files={"file": ("chain.xlsx", content)},
        headers=_HEADERS,
    )
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is True, body
    return str(body["dataset_id"])


def test_publish_chain_equivalence_and_dashboard_switch(chain_env) -> None:
    client, session = chain_env
    dataset_id = _upload(client, _workbook_bytes(with_income=True, with_bank=True), "接线账套")

    published = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS)
    assert published.status_code == 200, published.text
    receipt = published.json()
    assert receipt["published"] is True
    assert len(receipt["metric_snapshot_ids"]) == len(_PERIODS)
    assert receipt["analysis_run_status"] == "published"
    assert receipt["is_active_source"] is True

    # 口径等价：逐期间快照值 == 独立复算期望值（7 基线 + M1 结构 + M2 损益明细）
    for period, snapshot_id in zip(_PERIODS, receipt["metric_snapshot_ids"], strict=True):
        expected = _expected_period_metrics(period)
        detail_expected = _expected_detail_metrics(period)
        income_expected = _expected_income_metrics(period)
        cash_expected = _expected_cash_metrics(period)
        actual = _snapshot_values(session, snapshot_id)
        assert set(actual) == (
            set(expected) | set(detail_expected) | set(income_expected) | set(cash_expected)
        ), (period, set(actual))
        for code, expected_value in {
            **expected,
            **detail_expected,
            **income_expected,
            **cash_expected,
        }.items():
            assert actual[code] == expected_value, (period, code, actual[code])

    # 流量指标：全部期间快照之和 == derive_metrics 直读全期合计
    direct = derive_metrics(ImportStore(get_engine()), dataset_id)
    assert direct is not None
    direct_values = {item["metric_code"]: Decimal(item["value"]) for item in direct["metrics"]}
    summed = {
        code: sum(
            _snapshot_values(session, snapshot_id)[code]
            for snapshot_id in receipt["metric_snapshot_ids"]
        )
        for code in ("revenue", "total_expenses", "net_profit")
    }
    for code, value in summed.items():
        assert value == direct_values[code], (code, value, direct_values[code])

    # 治理切换：总览活跃源 = 本数据集最新快照；revenue 卡出真数，推不出的如实灰态
    # now 取数据窗口内（样例期间 2025 年），新鲜度判定才不会把样例判 stale
    overview = DashboardService().get_overview(
        session,
        filters=ActiveFilters(period_view="month", is_total_scope=True),
        now=datetime(2025, 3, 1, tzinfo=UTC),
    )
    # 两期数据凑不满 12 个月趋势面板 → degraded 是设计内如实降级（不补零）
    assert overview.state == "degraded"
    assert str(overview.context.metric_snapshot_id) == receipt["metric_snapshot_ids"][-1]
    expected_last = _expected_period_metrics("2025-02")
    revenue_card = next(card for card in overview.metric_cards if card.metric_code == "revenue")
    assert revenue_card.primary.status == "available"
    assert Decimal(str(revenue_card.primary.exact_value)) == expected_last["revenue"]
    # M2：利润表接线后毛利率卡出真数（P2 毛利率 4600/7000）
    margin_card = next(card for card in overview.metric_cards if card.metric_code == "gross_margin")
    assert margin_card.primary.status == "available"
    assert (
        Decimal(str(margin_card.primary.exact_value))
        == _expected_income_metrics("2025-02")["gross_margin"]
    )
    for code in ("orders", "operating_cash_flow"):
        card = next(c for c in overview.metric_cards if c.metric_code == code)
        assert card.primary.status == "unavailable", code


def test_publish_is_idempotent(chain_env) -> None:
    client, _session = chain_env
    dataset_id = _upload(client, _workbook_bytes(), "幂等账套")
    first = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS).json()
    second = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS).json()
    assert first["batch_id"] == second["batch_id"]
    assert first["analysis_run_id"] == second["analysis_run_id"]
    assert first["metric_snapshot_ids"] == second["metric_snapshot_ids"]

    # 无流水工作簿：现金口径指标缺席（不虚构）
    absent = _snapshot_values(_session, first["metric_snapshot_ids"][0])
    assert "cash_inflow" not in absent and "net_cash_flow" not in absent


def test_identity_mismatch_blocks_publish(chain_env) -> None:
    client, session = chain_env
    dataset_id = _upload(client, _identity_breaking_bytes(), "勾稽不平账套")
    response = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS)
    assert response.status_code == 422
    chain = client.get(f"/api/v1/importer/dataset/{dataset_id}/chain", headers=_HEADERS).json()
    assert chain["published"] is False
    assert (
        session.query(AnalysisBatch)
        .filter(AnalysisBatch.name == f"import-dataset-{dataset_id}")
        .count()
        == 0
    )


def test_publish_coexists_with_legacy_same_code_definition(chain_env) -> None:
    """回归（2026-10-29 真实环境暴露）：MetricDefinition 注册表按
    (metric_code, version) 全局唯一。库中已有同码旧定义（如 flow_v1 的
    revenue@1 或历史 TB 目录残留）时，TB 目录的备择定义走 version 2，
    发布不得 500，且旧定义内容不得被改动。"""
    client, session = chain_env
    from sqlalchemy import select

    from flow_api.infrastructure.models.analytics import MetricDefinition

    legacy = session.scalar(
        select(MetricDefinition).where(
            MetricDefinition.metric_code == "revenue",
            MetricDefinition.version == 1,
        )
    )
    legacy_payload = (
        None if legacy is None else (legacy.name, legacy.formula, legacy.aggregation, legacy.unit)
    )

    dataset_id = _upload(client, _workbook_bytes(with_income=True), "共存账套")
    receipt = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS)
    assert receipt.status_code == 200, receipt.text
    assert receipt.json()["published"] is True

    revenue_v2 = session.scalar(
        select(MetricDefinition).where(
            MetricDefinition.metric_code == "revenue",
            MetricDefinition.version == 2,
        )
    )
    assert revenue_v2 is not None
    # 发布后快照的 revenue 值挂在 v2 定义上（此前如存在同码 v1，其内容保持原样）
    if legacy_payload is not None:
        session.refresh(legacy)
        assert (
            legacy.name,
            legacy.formula,
            legacy.aggregation,
            legacy.unit,
        ) == legacy_payload


def test_bank_balance_discontinuity_rejected_at_upload(chain_env) -> None:
    """M3 上传门禁：流水余额不连续整批拒绝且逐行定位。"""
    client, _session = chain_env
    response = client.post(
        "/api/v1/importer/dataset?name=流水断档账套",
        files={
            "file": (
                "bank.xlsx",
                _workbook_bytes(with_bank=True, break_bank_balance=True),
            )
        },
        headers=_HEADERS,
    )
    body = response.json()
    assert body["ok"] is False
    assert any("余额不连续" in error["message"] for error in body["errors"])


def test_income_net_profit_mismatch_blocks_publish(chain_env) -> None:
    """M2 勾稽门禁：利润表净利润与余额表损益净额不符 → 发布拒绝且零残留。"""
    client, session = chain_env
    dataset_id = _upload(
        client,
        _workbook_bytes(with_income=True, net_override="9999"),
        "净利润不符账套",
    )
    response = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS)
    assert response.status_code == 422
    assert (
        session.query(AnalysisBatch)
        .filter(AnalysisBatch.name == f"import-dataset-{dataset_id}")
        .count()
        == 0
    )


def test_delete_withdraws_chain_and_dashboard_falls_back(chain_env) -> None:
    client, session = chain_env
    baseline_run_id = DashboardSourceRepository().get_latest(session).run.id
    dataset_id = _upload(client, _workbook_bytes(), "级联账套")
    receipt = client.post(f"/api/v1/importer/dataset/{dataset_id}/publish", headers=_HEADERS).json()
    assert receipt["published"] is True

    deleted = client.delete(f"/api/v1/importer/dataset/{dataset_id}", headers=_HEADERS)
    assert deleted.status_code == 200
    assert deleted.json()["chain_withdrawn"] is True

    # 数据集删除后 chain 端点 404（数据集不存在），发布链下线由 DB 断言证明
    gone = client.get(f"/api/v1/importer/dataset/{dataset_id}/chain", headers=_HEADERS)
    assert gone.status_code == 404

    assert (
        session.query(AnalysisBatch)
        .filter(AnalysisBatch.name == f"import-dataset-{dataset_id}")
        .count()
        == 0
    )
    assert session.get(MetricSnapshot, UUID(receipt["metric_snapshot_ids"][-1])) is None
    assert session.get(AnalysisRun, UUID(receipt["analysis_run_id"])) is None

    latest_run_id = DashboardSourceRepository().get_latest(session).run.id
    assert latest_run_id == baseline_run_id
