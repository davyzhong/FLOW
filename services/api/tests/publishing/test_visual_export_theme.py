from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from pptx import Presentation

from flow_api.visual_theme import apply_pptx_visual_theme, style_workbook_headers
from flow_api.visual_tokens import VISUAL_TOKENS


def test_print_report_uses_display_type_and_paper_gold_theme() -> None:
    from flow_api.statements.objective_report_html import _PRINT_CSS

    display_font = VISUAL_TOKENS["typography-family-display"].split(",", maxsplit=1)[0]
    assert display_font in _PRINT_CSS
    assert VISUAL_TOKENS["typography-family-numeric"].split(",", maxsplit=1)[0] in _PRINT_CSS
    assert "JetBrains Mono\", Menlo" not in _PRINT_CSS
    assert VISUAL_TOKENS["colors-brand-gold"] in _PRINT_CSS
    assert "@page" in _PRINT_CSS and "@media print" in _PRINT_CSS


def test_office_font_names_resolve_to_native_mac_families(monkeypatch) -> None:
    import flow_api.visual_theme as theme

    monkeypatch.setattr(theme.platform, "system", lambda: "Darwin")
    assert theme._primary_font("typography-family-display") == "Songti SC"
    assert theme._primary_font("typography-family-body") == "PingFang SC"


def test_pptx_has_paper_background_and_gold_heading_accent() -> None:
    import flow_api.visual_theme as theme

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "财务报告"

    apply_pptx_visual_theme(presentation)

    paper = VISUAL_TOKENS["colors-surface-canvas"].removeprefix("#")
    gold = VISUAL_TOKENS["colors-brand-gold"].removeprefix("#")
    assert str(slide.background.fill.fore_color.rgb) == paper
    title = slide.shapes.title
    assert (
        title.text_frame.paragraphs[0].font.name
        == theme._primary_font("typography-family-display")
    )
    assert str(slide.shapes[-1].fill.fore_color.rgb) == gold


def test_pptx_without_title_placeholder_still_styles_body_text() -> None:
    import flow_api.visual_theme as theme

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 3000000, 500000)
    textbox.text = "正文"

    apply_pptx_visual_theme(presentation)

    assert (
        textbox.text_frame.paragraphs[0].font.name
        == theme._primary_font("typography-family-body")
    )


def test_xlsx_theme_has_readable_print_layout() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["主题", "指标", "值", "单位", "期间", "口径", "来源"])
    sheet.append(["盈利能力", "毛利率", "0.25", "%", "FY2025", "收入减成本", "原始披露"])

    style_workbook_headers(workbook)
    restored = load_workbook(BytesIO(_save(workbook)))
    styled = restored.active
    assert styled is not None
    assert styled.page_setup.orientation == "landscape"
    assert styled.sheet_properties.pageSetUpPr.fitToPage is True
    assert styled.page_setup.fitToWidth == 1
    assert styled.print_title_rows == "$1:$1"
    assert styled.freeze_panes == "A2"
    assert styled.sheet_view.showGridLines is False

    evidence = restored.create_sheet("证据页")
    evidence.append(["长来源说明".ljust(70, "证")])
    style_workbook_headers(restored)
    assert restored["证据页"].column_dimensions["A"].width >= 84

    wide = restored.create_sheet("宽明细")
    for index in range(12):
        wide.cell(1, index + 1, f"列 {index + 1}")
    auto_width = restored.create_sheet("自动列宽")
    auto_width.append(["项目", "指标", "说明"])
    auto_width.append(["测试项", "数值", "说明".ljust(30, "明")])
    style_workbook_headers(restored)
    assert restored["宽明细"].page_setup.fitToWidth == 2
    assert restored["自动列宽"].column_dimensions["C"].width >= 32


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
