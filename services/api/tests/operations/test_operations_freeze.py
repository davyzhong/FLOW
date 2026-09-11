"""O3 slice-2：六主题概览冻结 + HTML/PDF 输出测试。

- 冻结走 U5 客观资格合同：draft 拒绝（report_not_published）、发布后通过；
- 幂等：同内容复用同版本；数据变更后重冻结版本递增；
- HTML 端点渲染事实卡片与范围说明；PDF 为真 Chromium 产物（不可用则跳过）。
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
import yaml
from alembic import command
from alembic.config import Config
from fastapi import HTTPException
from sqlalchemy import create_engine, delete
from sqlalchemy.orm import Session

from flow_api.infrastructure.models.statement import (
    StatementLineItem,
    StatementNormalizedItem,
    StatementReport,
)
from flow_api.operations.freeze import freeze_operations_overview
from flow_api.publishing.objective_freeze import freeze_objective_statement_report
from flow_api.settings import get_settings
from flow_api.statements.importer import import_statement_report
from flow_api.statements.normalization import normalize_report

REPO_ROOT = Path(__file__).resolve().parents[4]
SF_YAML = REPO_ROOT / "docs/implementation/p5/sf_2026q1_statements.yaml"


@pytest.fixture(scope="module", autouse=True)
def migrated_database() -> None:
    command.upgrade(Config("alembic.ini"), "head")


@pytest.fixture
def db_session() -> Iterator[Session]:
    engine = create_engine(get_settings().database_url)
    session = Session(engine, expire_on_commit=False)
    for table in (StatementNormalizedItem, StatementLineItem, StatementReport):
        session.execute(delete(table))
    session.commit()
    yield session
    session.close()
    engine.dispose()


def _seed(session: Session, *, status: str = "published") -> Any:
    payload: dict[str, Any] = yaml.safe_load(SF_YAML.read_text())
    report = import_statement_report(
        session,
        company_name="顺丰控股",
        stock_code="002352.SZ",
        report_kind="一季报",
        period_label="2026Q1",
        payload=payload,
        source_ref="p5_samples/sf_002352/SF_2026_Q1_report.pdf",
        source_sha256="a" * 64,
    )
    report.status = status
    session.flush()
    normalize_report(session, report)
    return report


def test_freeze_rejects_draft_then_succeeds_after_publish(
    db_session: Session,
) -> None:
    """U5 资格合同复用：未发布拒绝（typed code），发布后通过。"""
    report = _seed(db_session, status="draft")
    with pytest.raises(HTTPException) as error:
        # 直接调 freeze 服务层验证 typed code（路由层转换 HTTP 状态）
        from flow_api.publishing.objective_freeze import ObjectiveFreezeError

        try:
            freeze_operations_overview(db_session, report_id=report.id)
        except ObjectiveFreezeError as freeze_error:
            assert freeze_error.code == "report_not_published"
            raise HTTPException(status_code=422) from freeze_error
    assert error.value.status_code == 422

    report.status = "published"
    db_session.flush()
    snapshot = freeze_operations_overview(db_session, report_id=report.id)
    assert snapshot.report_type == "operations_overview"
    overview = snapshot.payload["overview"]
    assert len(overview["themes"]) == 6


def test_freeze_is_idempotent_and_versions_on_change(
    db_session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    from datetime import datetime as RealDateTime

    from flow_api.operations import freeze as freeze_module

    report = _seed(db_session)
    first = freeze_operations_overview(db_session, report_id=report.id)

    class LaterDateTime:
        @classmethod
        def now(cls) -> RealDateTime:
            return RealDateTime.fromisoformat("2030-01-02T03:04:05+08:00")

    monkeypatch.setattr(freeze_module, "datetime", LaterDateTime)
    second = freeze_operations_overview(db_session, report_id=report.id)
    assert second.id == first.id and second.version == first.version, "同内容幂等复用，不新建版本"

    # 数据变更（人工修正语义：直接改归一行值模拟）→ 重冻结版本递增
    row = (
        db_session.query(StatementNormalizedItem)
        .filter(
            StatementNormalizedItem.report_id == report.id,
            StatementNormalizedItem.item_id == "is.revenue",
        )
        .first()
    )
    row.value_current = row.value_current + 1
    db_session.flush()
    third = freeze_operations_overview(db_session, report_id=report.id)
    assert third.id != first.id and third.version == first.version + 1


def test_financial_and_operations_snapshots_have_independent_versions(
    db_session: Session,
) -> None:
    """同一财报的财务客观快照与经营快照都应能从 v1 开始并存。"""

    report = _seed(db_session)
    financial = freeze_objective_statement_report(db_session, report_id=report.id)
    operations = freeze_operations_overview(db_session, report_id=report.id)
    db_session.flush()

    assert financial.version == 1
    assert operations.version == 1
    assert financial.report_type == "objective_statement"
    assert operations.report_type == "operations_overview"


def test_overview_html_endpoint_renders_fact_cards(db_session: Session) -> None:
    from fastapi.testclient import TestClient

    from flow_api.api.routes.operations import get_operations_session
    from flow_api.main import create_app

    report = _seed(db_session)
    app = create_app()
    app.dependency_overrides[get_operations_session] = lambda: db_session
    client = TestClient(app)

    response = client.get(f"/api/v1/operations/overview/{report.id}/html")
    assert response.status_code == 200
    html = response.text
    assert "增长质量" in html and "盈利质量" in html
    assert "待内部数据" in html and "internal_data_required" in html
    assert "不在客观报告范围" in html

    freeze_response = client.post(f"/api/v1/operations/overview/{report.id}/freeze")
    assert freeze_response.status_code == 200
    body = freeze_response.json()
    assert body["report_type"] == "operations_overview"


def test_overview_xlsx_and_pptx_endpoints_use_frozen_payload(
    db_session: Session,
) -> None:
    import io

    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from pptx import Presentation

    from flow_api.api.routes.operations import get_operations_session
    from flow_api.main import create_app

    report = _seed(db_session)
    app = create_app()
    app.dependency_overrides[get_operations_session] = lambda: db_session
    client = TestClient(app)

    xlsx = client.get(f"/api/v1/operations/overview/{report.id}/xlsx")
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames[0] == "六主题概览"

    pptx = client.get(f"/api/v1/operations/overview/{report.id}/pptx")
    assert pptx.status_code == 200
    presentation = Presentation(io.BytesIO(pptx.content))
    assert presentation.slides[0].shapes.title.text == "经营分析概览"


def test_overview_pdf_is_real_chromium_output(db_session: Session, tmp_path: Path) -> None:
    from fastapi.testclient import TestClient

    from flow_api.api.routes.operations import get_operations_session
    from flow_api.main import create_app

    report = _seed(db_session)
    app = create_app()
    app.dependency_overrides[get_operations_session] = lambda: db_session
    client = TestClient(app)

    response = client.get(f"/api/v1/operations/overview/{report.id}/pdf")
    if response.status_code == 503:
        pytest.skip("Chromium 不可用，跳过真 PDF 验证")
    assert response.status_code == 200
    assert response.content[:5] == b"%PDF-", "必须是真实 PDF 魔数"
    assert b"/Page" in response.content or b"/Pages" in response.content
